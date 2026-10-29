#!/usr/bin/env python3
"""Read the comments reviewers left in an exported Word copy of a set.

    uv run --with python-docx --with pyyaml python3 prd_comments.py <set directory> <reviewed.docx>

Prints JSON: one entry per comment, in document order, with the id it is
anchored to, the text it marks, its author, date, and text, and its replies.
prd_export.py marks each section heading and each item with a bookmark named
by its id. A comment belongs to the nearest bookmark before the place it
starts. The set directory gives the ids, so each bookmark name maps back.

With --import it also adds each comment whose source the set does not hold yet
to <set>/17-review-comments.md, at status open. It creates and links that
document on the first import, and prints what it added and skipped.

    ... prd_comments.py <set directory> <reviewed.docx> --import
"""

import datetime
import json
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from prd_export import bookmark_name, read_set  # noqa: E402

try:
    import docx
    from docx.oxml.ns import qn
except ImportError:
    sys.exit("prd_comments: python-docx is required (uv run --with python-docx ...)")

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"


def ids_in(value, out):
    """Every id in a document's front matter, at any depth."""
    if isinstance(value, list):
        for one in value:
            ids_in(one, out)
    elif isinstance(value, dict):
        if isinstance(value.get("id"), str):
            out.append(value["id"])
        for key, one in value.items():
            if key != "links":
                ids_in(one, out)
    return out


def parents_of(path):
    """Each reply's parent, by the paraId of the comments' last paragraphs (commentsExtended.xml)."""
    with zipfile.ZipFile(path) as z:
        if "word/commentsExtended.xml" not in z.namelist():
            return {}
        from lxml import etree
        root = etree.fromstring(z.read("word/commentsExtended.xml"))
    return {ex.get(f"{{{W15}}}paraId"): ex.get(f"{{{W15}}}paraIdParent")
            for ex in root.iter(f"{{{W15}}}commentEx") if ex.get(f"{{{W15}}}paraIdParent")}


def read_comments(set_dir, path):
    root, rest = read_set(set_dir)
    names = {}
    for d in [root, *rest]:
        for ident in ids_in(d["front"], []):
            names.setdefault(bookmark_name(ident), ident)

    document = docx.Document(path)
    anchors, quotes, open_ranges = {}, {}, set()
    current = None
    for el in document.element.body.iter():
        tag = el.tag
        if tag == qn("w:bookmarkStart"):
            current = names.get(el.get(qn("w:name")), current)
        elif tag == qn("w:commentRangeStart"):
            cid = el.get(qn("w:id"))
            anchors[cid] = current
            quotes[cid] = []
            open_ranges.add(cid)
        elif tag == qn("w:commentRangeEnd"):
            open_ranges.discard(el.get(qn("w:id")))
        elif tag == qn("w:t") and el.text:
            for cid in open_ranges:
                quotes[cid].append(el.text)

    parents = parents_of(path)
    para_of, by_id = {}, {}
    for c in document.comments:
        last = c.paragraphs
        pid = last[-1]._p.get(f"{{{W14}}}paraId") if last else None
        cid = str(c.comment_id)
        if pid:
            para_of[pid] = cid
        by_id[cid] = {
            "source": f"{os.path.basename(path)}#{cid}",
            "about": anchors.get(cid),
            "quote": " ".join("".join(quotes.get(cid, [])).split()),
            "reviewer": c.author or c.initials or "unknown",
            "date": c.timestamp.date().isoformat() if c.timestamp else None,
            "comment": c.text.strip(),
            "replies": [],
        }
    top = []
    for c in document.comments:
        cid = str(c.comment_id)
        last = c.paragraphs
        pid = last[-1]._p.get(f"{{{W14}}}paraId") if last else None
        parent = para_of.get(parents[pid]) if pid in parents else None
        entry = by_id[cid]
        if parent is not None and parent in by_id:
            by_id[parent]["replies"].append({k: entry[k] for k in ("reviewer", "date", "comment")})
        else:
            top.append(entry)
    order = {cid: n for n, cid in enumerate(anchors)}
    top.sort(key=lambda e: order.get(e["source"].rsplit("#", 1)[1], len(order)))
    return top


REGISTER = "17-review-comments.md"


def block(key, text, indent):
    """A free-text field as a folded block scalar, which needs no escapes for any character."""
    return [f"{indent}{key}: >-", f"{indent}  {' '.join(str(text).split())}"]


def entry_lines(rc_id, c):
    out = [f"  - id: {rc_id}", f"    about: {c['about']}"]
    out += block("reviewer", c["reviewer"], "    ")
    if c.get("date"):
        out.append(f"    date: {c['date']}")
    if c.get("quote"):
        out += block("quote", c["quote"], "    ")
    out += block("comment", c["comment"], "    ")
    if c.get("replies"):
        out.append("    replies:")
        for r in c["replies"]:
            out += [f"      - reviewer: >-", f"          {' '.join(r['reviewer'].split())}"]
            if r.get("date"):
                out.append(f"        date: {r['date']}")
            out += block("comment", r["comment"], "        ")
    out += ["    status: open", f"    source: {c['source']}"]
    return out


def link_of(front, path):
    return f"{{ id: {front['id']}, kind: {front['kind']}, path: ./{path}, title: \"{front['title']}\" }}"


def import_comments(set_dir, path):
    """Adds new comments to the register, creating and linking it on the first import."""
    comments = read_comments(set_dir, path)
    root, rest = read_set(set_dir)
    register = os.path.join(set_dir, REGISTER)
    old = next((d for d in rest if d["front"]["kind"] == "review-comments"), None)
    known = {c.get("source") for c in (old["front"].get("comments") or [])} if old else set()
    fresh = [c for c in comments if c["source"] not in known and c["about"]]
    unplaced = [c for c in comments if c["source"] not in known and not c["about"]]
    taken = [int(c["id"][3:]) for c in (old["front"].get("comments") or [])] if old else []
    lines = []
    for n, c in enumerate(fresh, start=max(taken, default=0) + 1):
        lines += entry_lines(f"RC-{n:03d}", c)

    if old is None and fresh:
        last = max(rest, key=lambda d: d["front"].get("order", 0))
        me = {"id": f"{root['front']['id']}-review-comments", "kind": "review-comments", "title": "Review comments"}
        head = [
            "---", f"id: {me['id']}", "kind: review-comments", "schema: prd-forge/review-comments@1",
            "title: Review comments", "status: draft", f"updated: {datetime.date.today().isoformat()}",
            "section: 15", f"order: {last['front'].get('order', 160) + 10}", "comments:",
        ]
        tail = [
            "links:",
            f"  parent: {link_of(root['front'], root['file'])}",
            f"  prev: {link_of(last['front'], last['file'])}",
            "  next: null", "---", "", "# Review comments", "",
            "The comments reviewers left in an exported copy of the set, and the team's",
            "answer to each one. `/prd-comments` adds them, and the team answers them.", "",
        ]
        with open(register, "w", encoding="utf-8") as f:
            f.write("\n".join(head + lines + tail))
        # Link it: the last section's next, and the root's children.
        last_path = os.path.join(set_dir, last["file"])
        text = open(last_path, encoding="utf-8").read()
        text = re.sub(r"(?m)^  next: null$", f"  next: {link_of(me, REGISTER)}", text, count=1)
        open(last_path, "w", encoding="utf-8").write(text)
        root_path = os.path.join(set_dir, root["file"])
        text = open(root_path, encoding="utf-8").read()
        child = (f"    - {{ id: {me['id']}, kind: review-comments, path: ./{REGISTER}, "
                 f"title: Review comments, section: 15, order: {last['front'].get('order', 160) + 10}, status: draft }}")
        children = re.search(r"(?m)^  children:\n(?:    - .*\n)+", text)
        if children is None:
            sys.exit(f"prd_comments: {root['file']} has no links.children list to add {REGISTER} to")
        text = text[:children.end()] + child + "\n" + text[children.end():]
        open(root_path, "w", encoding="utf-8").write(text)
    elif fresh:
        text = open(register, encoding="utf-8").read()
        at = re.search(r"(?m)^links:$", text)
        if at is None:
            sys.exit(f"prd_comments: {REGISTER} has no links block")
        text = text[:at.start()] + "\n".join(lines) + "\n" + text[at.start():]
        text = re.sub(r"(?m)^comments: \[\]$", "comments:", text, count=1)
        open(register, "w", encoding="utf-8").write(text)
    return {
        "added": [c["source"] for c in fresh],
        "skipped": [c["source"] for c in comments if c["source"] in known],
        "unplaced": unplaced,
    }


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    if len(args) != 2:
        sys.exit("usage: prd_comments.py <set directory> <reviewed.docx> [--import]")
    if "--import" in argv:
        print(json.dumps(import_comments(args[0], args[1]), indent=2, ensure_ascii=False))
    else:
        print(json.dumps(read_comments(args[0], args[1]), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1:])
